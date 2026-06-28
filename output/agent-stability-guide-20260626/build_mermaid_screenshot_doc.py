from __future__ import annotations

import json
import os
import subprocess
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "agent-stability-guide-20260626"
ASSET_DIR = OUT / "mermaid-screenshots"
DOCX_PATH = OUT / "零部件设计智能体稳定性与正确率工作指导书-Mermaid截图版.docx"
ARCH_DOCX_PATH = OUT / "零部件设计智能体软件模型架构说明书-Mermaid截图优化版.docx"
OPS_DOCX_PATH = OUT / "零部件设计智能体持续改进操作说明书-Mermaid截图优化版.docx"
PNPM = Path(r"C:\Users\ASUS\.cache\codex-runtimes\codex-primary-runtime\dependencies\bin\pnpm.cmd")


MERMAID_INIT = """%%{init: {"theme": "base", "themeVariables": {"fontFamily": "Microsoft YaHei, Segoe UI, Arial", "primaryColor": "#EAF3FF", "primaryTextColor": "#1F2937", "primaryBorderColor": "#2F5F98", "lineColor": "#4B5563", "secondaryColor": "#EEF7F1", "tertiaryColor": "#FFF7E6", "background": "#FFFFFF", "mainBkg": "#FFFFFF", "clusterBkg": "#F8FAFC", "clusterBorder": "#CBD5E1", "noteBkgColor": "#FFF7E6", "noteTextColor": "#111827", "actorBkg": "#EAF3FF", "actorBorder": "#2F5F98", "actorTextColor": "#111827", "activationBkgColor": "#DBEAFE", "activationBorderColor": "#2F5F98"}, "flowchart": {"htmlLabels": true, "curve": "basis", "nodeSpacing": 52, "rankSpacing": 64}, "sequence": {"mirrorActors": false, "showSequenceNumbers": true, "wrap": true, "width": 180}} }%%"""


STYLE = """
classDef main fill:#EAF3FF,stroke:#2F5F98,stroke-width:2px,color:#111827;
classDef agent fill:#EEF7F1,stroke:#2F7D4F,stroke-width:2px,color:#111827;
classDef tool fill:#FFF7E6,stroke:#B7791F,stroke-width:2px,color:#111827;
classDef data fill:#F5F3FF,stroke:#6D5BD0,stroke-width:2px,color:#111827;
classDef risk fill:#FEE2E2,stroke:#B91C1C,stroke-width:2px,color:#111827;
classDef out fill:#ECFEFF,stroke:#0E7490,stroke-width:2px,color:#111827;
"""


@dataclass(frozen=True)
class AgentSpec:
    key: str
    title: str
    agent_file: str
    responsibility: str
    skills: str
    tools: str
    outputs: str
    failure_signals: str
    no_code_handles: str
    architecture: str
    workflow: str


def flow_init(body: str) -> str:
    return f"{MERMAID_INIT}\nflowchart TB\n{body}\n{STYLE}\n"


def seq_init(body: str) -> str:
    return f"{MERMAID_INIT}\nsequenceDiagram\n{body}\n"


OVERALL_ARCH = flow_init(
    r"""
    U["用户 / 业务工程师<br/>任务、确认、补充材料"]:::main
    AI["AI Service<br/>call_agent / stream_turn"]:::main
    BR["RuntimeBridgeManager<br/>turn-stream-v1 协议桥接"]:::main
    APP["本地 ClientApp<br/>/mcp/&lt;toolbox&gt; 工具入口"]:::main
    LIVE["TJUAE LiveCoordinator<br/>会话、插件、工具定义同步"]:::main
    DA["design-agent<br/>只做任务分派与 AskUserQuestion"]:::agent
    A1["task-query-display-agent<br/>任务抓取与展示"]:::agent
    A2["nx-template-agent<br/>模板调用与工作部件确认"]:::agent
    A3["parameter-intelligence-agent<br/>证据、公式、参数写入计划"]:::agent
    A4["nx-modeling-tc-agent<br/>NX参数执行与TC回传"]:::agent
    A5["drawing-update-agent<br/>二维图纸与视图证据"]:::agent
    A6["dfmea-author-agent<br/>DFMEA模板化编写"]:::agent
    A7["design-report-author-agent<br/>设计说明书生成"]:::agent
    MCP["本地 MCP 工具箱<br/>connector / nx / teamcenter / dfmea / report / local_file"]:::tool
    SYS["企业系统<br/>IPM / QPP / ECR / MySQL / NX / Teamcenter"]:::data
    ART["交付物<br/>模型、截图、DFMEA、说明书"]:::out
    LOG["conversation JSONL + summary<br/>tool_events / last_error / artifact"]:::data
    WB["评测工作台<br/>7任务 35场景 134检查点"]:::data
    IMP["其他智能体改进闭环<br/>日志证据 / 指标分析 / 失败归因 / 指令修订 / 复测"]:::risk

    U --> AI --> BR --> APP --> LIVE --> DA
    DA --> A1
    DA --> A2
    DA --> A3
    DA --> A4
    DA --> A5
    DA --> A6
    DA --> A7
    A1 --> MCP
    A2 --> MCP
    A3 --> MCP
    A4 --> MCP
    A5 --> MCP
    A6 --> MCP
    A7 --> MCP
    MCP --> SYS
    MCP --> ART
    APP --> LOG
    LOG --> IMP
    WB --> IMP
    IMP --> DA
    IMP --> A1
    IMP --> A2
    IMP --> A3
    IMP --> A4
    IMP --> A5
    IMP --> A6
    IMP --> A7
"""
)

IMPROVEMENT_LOOP = flow_init(
    r"""
    F["发现失败<br/>用户反馈或评测记录未通过"]:::risk
    L["日志证据智能体<br/>读取 JSONL / summary<br/>还原工具调用与错误"]:::agent
    M["指标分析智能体<br/>引用评测工作台<br/>定位任务-场景-检查点-版本差异"]:::agent
    R["失败归因智能体<br/>归类到路由、边界、证据、模板、外部系统、代码缺陷"]:::agent
    P["指令/模板修订智能体<br/>只提出无代码修订<br/>Agent说明、Skill、模板、SOP、测试样例"]:::agent
    H["人工评审<br/>确认不破坏边界与事实来源"]:::main
    T["复测智能体<br/>固定原失败场景 + 相邻场景<br/>记录版本对比"]:::agent
    G{"连续通过<br/>且无新增阻断?"}:::data
    CLOSE["归档改进记录<br/>失败样本、修改点、复测结果"]:::out
    ESC["升级缺陷<br/>仅当证据指向 MCP / NX / TC / 桥接代码故障"]:::risk

    F --> L --> M --> R --> P --> H --> T --> G
    G -- 是 --> CLOSE
    G -- 否：继续无代码修订 --> L
    R -- 明确代码故障 --> ESC
"""
)


AGENTS: list[AgentSpec] = [
    AgentSpec(
        key="task_query",
        title="子智能体1：任务抓取与展示",
        agent_file="task-query-display-agent.md",
        responsibility="从 IPM/QPP/ECR 等系统读取候选任务，保留原始任务内容，向主智能体返回完整任务快照。",
        skills="enterprise-task-intake",
        tools="query_ipm_list、connect_qpp、query_ecr_list、local_python_task_start/status、Read、Write",
        outputs="任务候选清单、被选任务快照、系统来源、原始字段、缺失字段提示。",
        failure_signals="候选任务丢失；字段被改写；来源系统混淆；慢查询超时后没有可解释提示。",
        no_code_handles="补充任务筛选规则示例；要求输出原始字段和来源；为慢查询增加先返回候选范围的交互话术；将失败样本录入任务1评测场景。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>调用任务抓取子智能体"]:::main
    A["task-query-display-agent<br/>只负责任务查询与展示"]:::agent
    S["Skill: enterprise-task-intake<br/>企业任务接入流程"]:::data
    T1["IPM工具<br/>query_ipm_list"]:::tool
    T2["QPP工具<br/>connect_qpp"]:::tool
    T3["ECR工具<br/>query_ecr_list"]:::tool
    PY["local_python_task<br/>长查询异步状态"]:::tool
    RAW["原始任务字段<br/>来源、编号、描述、附件线索"]:::data
    SNAP["完整任务快照<br/>不擅自改写业务事实"]:::out
    LOG["日志与评测<br/>任务1 5场景 19检查点"]:::data
    IMP["其他智能体<br/>日志证据 + 指标归因"]:::risk

    M --> A --> S
    S --> T1 --> RAW
    S --> T2 --> RAW
    S --> T3 --> RAW
    S --> PY --> RAW
    RAW --> SNAP --> M
    A --> LOG --> IMP --> S
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as task-query-display-agent
    participant Q as IPM/QPP/ECR工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提交任务关键词、编号或用户条件
    A->>Q: 查询候选任务并保留来源系统字段
    Q-->>A: 返回候选列表、原始描述、附件线索
    A->>A: 去重、标记缺失字段、禁止改写任务事实
    A-->>M: 返回候选清单和被选任务完整快照
    A->>L: 写入工具输入、输出、耗时、异常
    L->>W: 失败时绑定任务1的场景和检查点
    Note over A,W: 修订抓手是筛选话术、字段保留要求、慢查询处理和评测样本，不先改代码
"""
        ),
    ),
    AgentSpec(
        key="template",
        title="子智能体2：NX零部件模板调用",
        agent_file="nx-template-agent.md",
        responsibility="依据任务类型选择 Teamcenter/NX 模板，打开或新建零部件，并确认 NX WorkPart。",
        skills="template-source-selection、nx-workpart-confirmation",
        tools="Teamcenter folder/classification/children、nx_open_tcpart、nx_open_part、nx_get_work_part_info、nx_create_new_part",
        outputs="模板来源、零件打开结果、WorkPart确认信息、新建部件标识。",
        failure_signals="模板选错；未确认 WorkPart；打开了图纸而不是模型；Teamcenter路径或分类口径不一致。",
        no_code_handles="补充模板选择优先级和反例；明确必须回读 WorkPart；将路径/分类别名放入模板选择知识表；用评测任务2复测。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>提供任务快照和模板需求"]:::main
    A["nx-template-agent<br/>模板选择 + WorkPart确认"]:::agent
    S1["Skill: template-source-selection<br/>模板来源选择"]:::data
    S2["Skill: nx-workpart-confirmation<br/>工作部件确认"]:::data
    TC["Teamcenter工具<br/>文件夹、分类、子项"]:::tool
    NX1["NX打开工具<br/>nx_open_tcpart / nx_open_part"]:::tool
    NX2["NX创建与回读<br/>nx_create_new_part / nx_get_work_part_info"]:::tool
    SRC["候选模板<br/>TC对象、分类、路径、版本"]:::data
    WP["WorkPart确认结果<br/>部件号、名称、单位、状态"]:::out
    LOG["日志与评测<br/>任务2 5场景 22检查点"]:::data
    IMP["其他智能体<br/>模板误选归因与指令修订"]:::risk

    M --> A --> S1 --> TC --> SRC
    SRC --> NX1 --> NX2 --> WP --> M
    A --> S2 --> NX2
    A --> LOG --> IMP --> S1
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as nx-template-agent
    participant TC as Teamcenter工具
    participant NX as NX工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供任务类型、部件族、模板约束
    A->>TC: 查找文件夹、分类、子项候选
    TC-->>A: 返回模板对象、版本、路径
    A->>A: 按模板选择规则排序并说明依据
    A->>NX: 打开TC模板或创建新零件
    NX-->>A: 返回打开状态
    A->>NX: 回读 nx_get_work_part_info
    NX-->>A: 返回当前 WorkPart 信息
    A-->>M: 返回模板来源和 WorkPart 确认结果
    A->>L: 记录TC/NX调用链和失败点
    L->>W: 失败时关联任务2检查点
    Note over A,W: 重点修订模板选择规则、别名表和WorkPart回读门禁
"""
        ),
    ),
    AgentSpec(
        key="parameters",
        title="子智能体3：参数智能化处理",
        agent_file="parameter-intelligence-agent.md",
        responsibility="抽取任务证据、使用权威公式计算参数，匹配 NX 表达式并产出参数写入计划。",
        skills="parameter-evidence-extraction、formula-parameter-calculation、parameter-write-plan",
        tools="mysql_query、nx_get_work_part_info、nx_get_all_params_list、nx_get_drive_params_list、nx_find_params、nx_resolve_parameter、local_python_task_start/status、Read、Write",
        outputs="参数证据表、公式计算过程、表达式映射、写入计划、风险和人工确认项。",
        failure_signals="套用错误公式；单位换算错误；参数名匹配错；证据缺失但仍计算；把计划误写入 NX。",
        no_code_handles="把 Excel 公式来源作为唯一权威；新增单位换算检查点；参数同义词表只作为候选需回读确认；评测任务3覆盖边界值和缺证场景。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>提供任务快照、模板和用户确认"]:::main
    A["parameter-intelligence-agent<br/>证据、公式、映射、写入计划"]:::agent
    E["Skill: evidence-extraction<br/>字段、附件、用户输入抽取"]:::data
    F["Skill: formula-calculation<br/>Excel公式权威计算"]:::data
    P["Skill: write-plan<br/>仅生成计划不写NX"]:::data
    DB["mysql_query<br/>公式、参数、历史样例"]:::tool
    NX["NX只读工具<br/>WorkPart / all params / drive params / resolve"]:::tool
    PY["local_python_task<br/>公式校验脚本"]:::tool
    MAP["参数映射<br/>NX表达式、单位、来源证据"]:::data
    PLAN["参数写入计划<br/>待执行值、置信度、确认项"]:::out
    LOG["日志与评测<br/>任务3 5场景 20检查点"]:::data
    IMP["其他智能体<br/>公式/单位/映射失败归因"]:::risk

    M --> A --> E --> DB
    A --> F --> PY
    A --> P --> NX --> MAP
    DB --> MAP
    PY --> MAP --> PLAN --> M
    A --> LOG --> IMP --> E
    IMP --> F
    IMP --> P
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as parameter-intelligence-agent
    participant DB as MySQL/资料源
    participant NX as NX只读工具
    participant PY as 公式校验脚本
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供任务快照、模板信息、用户补充
    A->>DB: 查询公式来源、参数说明、历史样例
    A->>NX: 读取WorkPart和表达式列表
    A->>A: 抽取证据并标记缺失项
    A->>PY: 用权威公式计算并校验单位
    PY-->>A: 返回计算结果和校验记录
    A->>NX: 解析候选表达式，禁止写入
    NX-->>A: 返回表达式匹配和单位信息
    A-->>M: 输出参数写入计划和需确认项
    A->>L: 记录证据、公式、脚本结果、映射依据
    L->>W: 失败时关联任务3检查点
    Note over A,W: 修订抓手是公式证据、单位检查、参数同义词和缺证阻断规则
"""
        ),
    ),
    AgentSpec(
        key="modeling",
        title="子智能体4：参数化建模及TC回传",
        agent_file="nx-modeling-tc-agent.md",
        responsibility="执行已确认的参数写入计划，回读 NX 参数与模型状态，保存并回传 Teamcenter。",
        skills="nx-parameter-execution、tc-release-preparation",
        tools="NX read/resolve/update/batch/create/save、Teamcenter upload/copy/export/personnel、local_file_exists/read_bytes",
        outputs="写入前后对比、NX保存结果、Teamcenter回传对象、失败回滚/阻断说明。",
        failure_signals="未按计划写入；未回读验证；保存失败未阻断；Teamcenter上传对象错误；把计算职责混入执行阶段。",
        no_code_handles="要求写前计划哈希、写后回读表、失败即停止；把TC回传清单模板化；用评测任务4验证写入、保存和回传。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>交付已确认参数计划"]:::main
    A["nx-modeling-tc-agent<br/>NX执行 + TC回传"]:::agent
    S1["Skill: nx-parameter-execution<br/>写入、更新、回读"]:::data
    S2["Skill: tc-release-preparation<br/>回传准备"]:::data
    NX["NX写入工具<br/>resolve / update / batch / create / save"]:::tool
    TC["Teamcenter工具<br/>upload / copy / export / personnel"]:::tool
    FILE["local_file<br/>存在性与字节读取"]:::tool
    PLAN["已确认参数计划<br/>值、单位、表达式、签名"]:::data
    READBACK["写后回读证据<br/>表达式值、模型状态、保存状态"]:::data
    UPLOAD["TC回传结果<br/>对象、版本、附件、人员"]:::out
    LOG["日志与评测<br/>任务4 5场景 19检查点"]:::data
    IMP["其他智能体<br/>执行链归因与门禁修订"]:::risk

    M --> PLAN --> A --> S1 --> NX --> READBACK
    READBACK --> S2 --> TC --> UPLOAD --> M
    FILE --> TC
    A --> LOG --> IMP --> S1
    IMP --> S2
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as nx-modeling-tc-agent
    participant NX as NX写入工具
    participant TC as Teamcenter工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供已确认参数写入计划
    A->>A: 校验计划完整性和人工确认标记
    A->>NX: 解析表达式并执行批量写入
    NX-->>A: 返回写入状态
    A->>NX: 更新模型、保存、回读表达式和WorkPart
    NX-->>A: 返回写后证据和保存状态
    A->>TC: 上传或复制相关对象并准备回传
    TC-->>A: 返回TC对象、版本、附件结果
    A-->>M: 返回写入前后对比和TC回传清单
    A->>L: 记录计划签名、写入结果、回读差异
    L->>W: 失败时关联任务4检查点
    Note over A,W: 修订抓手是写前门禁、写后回读表、保存失败阻断、TC回传清单
"""
        ),
    ),
    AgentSpec(
        key="drawing",
        title="子智能体5：二维图纸自动绘制",
        agent_file="drawing-update-agent.md",
        responsibility="选择 Teamcenter 图纸/规范，打开 NX 图纸视图并生成截图证据，不承担建模和DFMEA职责。",
        skills="tc-drawing-selection、nx-view-screenshot-evidence",
        tools="TC children/export、NX open TC drawing、get views、switch、fit、create image、get/set style、local_file_exists",
        outputs="图纸对象、视图清单、样式状态、截图证据、缺失视图/规范说明。",
        failure_signals="选错图纸版本；视图为空；截图不清晰或没有视图上下文；样式变更未还原；图纸缺失未向用户确认。",
        no_code_handles="图纸选择优先级模板化；截图命名和视图清单固定输出；新增空视图阻断话术；用评测任务5补齐未测与图纸边界场景。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>提供模型/TC对象/图纸需求"]:::main
    A["drawing-update-agent<br/>图纸选择 + 视图截图证据"]:::agent
    S1["Skill: tc-drawing-selection<br/>TC图纸/规范选择"]:::data
    S2["Skill: nx-view-screenshot-evidence<br/>视图切换与截图"]:::data
    TC["Teamcenter工具<br/>children / export"]:::tool
    NX["NX图纸工具<br/>open drawing / get views / switch / fit / image / style"]:::tool
    FS["local_file_exists<br/>截图和导出文件检查"]:::tool
    D["图纸候选<br/>对象、版本、规范、来源"]:::data
    V["视图证据<br/>视图清单、截图、样式状态"]:::out
    LOG["日志与评测<br/>任务5 5场景 19检查点"]:::data
    IMP["其他智能体<br/>图纸选择与截图质量归因"]:::risk

    M --> A --> S1 --> TC --> D
    D --> S2 --> NX --> V --> FS --> M
    A --> LOG --> IMP --> S1
    IMP --> S2
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as drawing-update-agent
    participant TC as Teamcenter工具
    participant NX as NX图纸工具
    participant FS as local_file工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供TC对象和图纸/规范需求
    A->>TC: 查询子项并导出候选图纸或规范
    TC-->>A: 返回图纸对象、版本、路径
    A->>NX: 打开TC图纸并读取视图列表
    NX-->>A: 返回视图名称和状态
    A->>NX: 切换视图、fit、生成截图、检查样式
    NX-->>A: 返回截图路径和样式信息
    A->>FS: 检查截图文件存在
    FS-->>A: 返回文件状态
    A-->>M: 输出图纸对象、视图清单、截图证据
    A->>L: 记录图纸选择、视图切换、截图路径
    L->>W: 失败时关联任务5检查点
    Note over A,W: 修订抓手是图纸选择优先级、空视图阻断、截图命名和样式恢复
"""
        ),
    ),
    AgentSpec(
        key="dfmea",
        title="子智能体6：DFMEA自动编写",
        agent_file="dfmea-author-agent.md",
        responsibility="基于受控模板和风险规则生成 DFMEA 工作簿，并通过模板校验。",
        skills="dfmea-risk-authoring、dfmea-workbook-generation",
        tools="TC children/export、dfmea_template_list/inspect/select/calculate_risk/fill_template/validate_workbook、local_file_exists/read_bytes",
        outputs="DFMEA XLSX、风险计算记录、模板校验结果、缺失证据和人工确认项。",
        failure_signals="不用受控模板；风险分值不合规；字段空缺未阻断；模板校验失败仍交付；事实来源不明。",
        no_code_handles="把DFMEA字段证据表固定化；风险分值必须走calculate_risk；模板校验失败不得交付；针对当前低分任务6优先扩充失败样本和修订生成要求。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>提供任务、模型、参数、图纸证据"]:::main
    A["dfmea-author-agent<br/>风险编写 + 工作簿生成"]:::agent
    S1["Skill: dfmea-risk-authoring<br/>失效模式、影响、原因、控制"]:::data
    S2["Skill: dfmea-workbook-generation<br/>受控模板填充与校验"]:::data
    TC["Teamcenter工具<br/>children / export"]:::tool
    DF["DFMEA工具<br/>template_list / inspect / select / calculate_risk / fill / validate"]:::tool
    FS["local_file<br/>存在性与字节读取"]:::tool
    EV["输入证据<br/>设计参数、图纸、任务约束、历史风险"]:::data
    XLSX["DFMEA工作簿<br/>风险记录 + 校验结果"]:::out
    LOG["日志与评测<br/>任务6 5场景 17检查点"]:::data
    IMP["其他智能体<br/>低分项优先归因与修订"]:::risk

    M --> EV --> A --> S1 --> TC
    A --> S2 --> DF --> XLSX --> FS --> M
    A --> LOG --> IMP --> S1
    IMP --> S2
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as dfmea-author-agent
    participant TC as Teamcenter工具
    participant DF as DFMEA工具
    participant FS as local_file工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供任务、设计证据、图纸/参数结果
    A->>TC: 获取必要附件和历史对象
    TC-->>A: 返回证据文件和对象
    A->>DF: 列出、检查并选择受控模板
    DF-->>A: 返回模板结构和必填字段
    A->>DF: 计算风险值并填充模板
    DF-->>A: 返回DFMEA工作簿路径
    A->>DF: validate_workbook校验
    DF-->>A: 返回校验结果
    A->>FS: 检查工作簿文件
    A-->>M: 仅在校验通过后返回DFMEA交付物
    A->>L: 记录模板、风险计算、校验错误
    L->>W: 失败时关联任务6检查点
    Note over A,W: 修订抓手是证据字段表、风险计算规则、模板校验阻断和高频失败样本
"""
        ),
    ),
    AgentSpec(
        key="report",
        title="子智能体7：设计说明书自动编写",
        agent_file="design-report-author-agent.md",
        responsibility="汇总任务、参数、模型、图纸、DFMEA等证据，按设计说明书模板生成 DOCX。",
        skills="report-evidence-collection、report-docx-generation",
        tools="design_report_template_inspect/start/collect/decrypt_image/answer/status/generate/update/get_artifact、local_file_exists/read_bytes",
        outputs="设计说明书 DOCX、证据引用清单、图片解密/引用结果、待用户确认项。",
        failure_signals="引用了未确认事实；图片缺失或打不开；模板章节缺失；生成状态未完成就交付；报告与模型/DFMEA不一致。",
        no_code_handles="固定证据清单和章节映射；图片必须经过decrypt/check；缺证章节输出待确认项；用评测任务7做一致性回归。",
        architecture=flow_init(
            r"""
    M["design-agent<br/>汇总前序交付物和用户确认"]:::main
    A["design-report-author-agent<br/>证据收集 + DOCX生成"]:::agent
    S1["Skill: report-evidence-collection<br/>事实来源和章节证据"]:::data
    S2["Skill: report-docx-generation<br/>模板生成和状态检查"]:::data
    RP["设计报告工具<br/>template_inspect / start / collect / decrypt_image / answer / generate / status / artifact"]:::tool
    FS["local_file<br/>存在性与字节读取"]:::tool
    EV["证据包<br/>任务、参数、模型、截图、DFMEA、确认项"]:::data
    DOC["设计说明书DOCX<br/>章节、图片、引用、状态"]:::out
    LOG["日志与评测<br/>任务7 5场景 18检查点"]:::data
    IMP["其他智能体<br/>事实一致性与模板缺口归因"]:::risk

    M --> EV --> A --> S1 --> RP
    A --> S2 --> RP --> DOC --> FS --> M
    A --> LOG --> IMP --> S1
    IMP --> S2
"""
        ),
        workflow=seq_init(
            r"""
    participant M as design-agent
    participant A as design-report-author-agent
    participant RP as 设计报告工具
    participant FS as local_file工具
    participant L as 日志证据
    participant W as 评测工作台

    M->>A: 提供任务、参数、模型、图纸、DFMEA证据
    A->>RP: inspect模板并启动报告任务
    RP-->>A: 返回模板章节和任务ID
    A->>RP: collect证据、decrypt_image、answer缺口问题
    RP-->>A: 返回证据接收状态
    A->>RP: generate并轮询status
    RP-->>A: 返回生成完成状态和artifact
    A->>FS: 检查DOCX存在并读取元信息
    A-->>M: 返回设计说明书和证据引用清单
    A->>L: 记录模板、证据、图片、生成状态
    L->>W: 失败时关联任务7检查点
    Note over A,W: 修订抓手是章节证据表、图片检查、生成状态门禁和一致性复测
"""
        ),
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 100, start: int = 140, bottom: int = 100, end: int = 140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        tag = qn(f"w:{side}")
        for child in list(tc_mar):
            if child.tag == tag:
                tc_mar.remove(child)
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def set_table_borders(table, color: str = "B8C4D1") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        for child in list(borders):
            if child.tag == tag:
                borders.remove(child)
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)


def set_run_family(run, east_asia: str = "Microsoft YaHei", latin: str = "Aptos") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    set_run_family(r)
    r.font.size = Pt(8.6 if not bold else 8.8)
    r.font.color.rgb = RGBColor(31, 41, 55)


def add_hyperlink_like_path(doc: Document, path: Path) -> None:
    p = doc.add_paragraph()
    p.style = "Body Text"
    run = p.add_run(str(path))
    set_run_family(run, east_asia="Microsoft YaHei", latin="Consolas")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(31, 78, 121)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    styles = doc.styles
    for style_name in ["Normal", "Body Text", "List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor(31, 41, 55)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.18

    for name, size, color in [
        ("Title", 21, "17365D"),
        ("Heading 1", 16, "17365D"),
        ("Heading 2", 13, "1F4E79"),
        ("Heading 3", 11.2, "365F91"),
    ]:
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name == "Heading 3" else 12)
        style.paragraph_format.space_after = Pt(5 if name == "Heading 3" else 7)
        style.paragraph_format.line_spacing = 1.12

    return doc


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("零部件设计智能体稳定性与正确率工作指导书")
    r.bold = True
    r.font.size = Pt(21)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.color.rgb = RGBColor(23, 54, 93)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Mermaid截图版 | 面向7个子智能体的错误处理、评测闭环与无代码改进操作手册")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(80, 80, 80)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["版本", "适用对象", "依据", "生成日期"]
    values = ["Mermaid截图版", "业务工程师 / 智能体维护工程师", "项目代码 + 现有评测工作台 + 运行日志", "2026-06-26"]
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
        set_cell_text(table.rows[0].cells[i], f"{h}\n{values[i]}", bold=i == 0)


def add_title_page(
    doc: Document,
    *,
    title: str,
    subtitle: str,
    version: str,
    audience: str,
    basis: str,
    doc_type: str,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    set_run_family(r)
    r.font.color.rgb = RGBColor(23, 54, 93)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    set_run_family(r)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(80, 80, 80)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table, "AAB7C4")
    headers = ["文档类型", "版本", "适用对象", "依据", "生成日期"]
    values = [doc_type, version, audience, basis, "2026-06-26"]
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "EEF3F8")
        set_cell_text(table.rows[0].cells[i], f"{h}\n{values[i]}", bold=i == 0)


def add_para(doc: Document, text: str, style: str = "Body Text") -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    set_run_family(r)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(31, 41, 55)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.16
        r = p.add_run(item)
        set_run_family(r)
        r.font.size = Pt(9.6)


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(48, 48, 48)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.16
        r = p.add_run(item)
        set_run_family(r)
        r.font.size = Pt(9.6)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "EEF3F8")
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
                set_cell_margins(row.cells[i])


def add_diagram(doc: Document, title: str, image_path: Path, caption: str) -> None:
    doc.add_page_break()
    doc.add_heading(title, level=3)
    with Image.open(image_path) as image:
        px_w, px_h = image.size
    max_w_cm = 24.7
    max_h_cm = 16.0
    scale = min(max_w_cm / px_w, max_h_cm / px_h)
    width_cm = px_w * scale
    height_cm = px_h * scale
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(90, 90, 90)


def write_mermaid_sources(diagrams: dict[str, str]) -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for old in ASSET_DIR.glob("*"):
        if old.is_file():
            old.unlink()
    for key, source in diagrams.items():
        (ASSET_DIR / f"{key}.mmd").write_text(source, encoding="utf-8")
    browser_candidates = [
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    ]
    browser = next((path for path in browser_candidates if path.exists()), None)
    config = {
        "args": [
            "--no-sandbox",
            "--disable-setuid-sandbox",
            "--disable-dev-shm-usage",
        ]
    }
    if browser:
        config["executablePath"] = str(browser)
    (ASSET_DIR / "puppeteer-config.json").write_text(json.dumps(config), encoding="utf-8")


def render_diagrams(diagrams: dict[str, str]) -> dict[str, Path]:
    rendered: dict[str, Path] = {}
    for key in diagrams:
        src = ASSET_DIR / f"{key}.mmd"
        out = ASSET_DIR / f"{key}.png"
        cmd = [
            str(PNPM),
            "dlx",
            "@mermaid-js/mermaid-cli",
            "-i",
            str(src),
            "-o",
            str(out),
            "-p",
            str(ASSET_DIR / "puppeteer-config.json"),
            "-w",
            "1800",
            "-H",
            "1150",
            "-b",
            "white",
            "-s",
            "2",
        ]
        subprocess.run(cmd, check=True, cwd=str(ROOT))
        rendered[key] = out
    return rendered


def build_doc(rendered: dict[str, Path]) -> None:
    doc = setup_document()
    add_title(doc)

    doc.add_heading("1. 使用边界", level=1)
    add_para(
        doc,
        "本指导书面向 mc-design 项目的7个任务子智能体。目标不是重新发明评测指标，也不是直接进入代码开发，而是让工程师在后续使用中，依靠现有日志、现有评测工作台和其他智能体完成可追踪、可复测、可落地的稳定性与正确率改进。",
    )
    add_bullets(
        doc,
        [
            "主智能体 design-agent 只负责路由、编排和必要的人机确认；具体任务由7个子智能体承担。",
            "常规改进对象是 Agent 指令、Skill 流程、输出模板、知识证据、评测样本、人工确认话术和运行SOP；只有证据明确指向 MCP/NX/TC/桥接程序缺陷时才进入代码缺陷流程。",
            "错误处理必须留下三类证据：原始日志、评测检查点、修订前后复测结果。没有这三类证据，不关闭问题。",
        ],
    )

    doc.add_heading("2. 总体架构", level=1)
    add_diagram(
        doc,
        "总体架构图",
        rendered["overall_arch"],
        "图1：运行链路、7个子智能体、企业系统、日志与评测改进闭环。",
    )
    add_diagram(
        doc,
        "稳定性与正确率改进闭环",
        rendered["improvement_loop"],
        "图2：由其他智能体处理日志和指标，形成无代码修订与复测闭环。",
    )

    doc.add_heading("3. 七个子智能体边界速查", level=1)
    add_table(
        doc,
        ["任务", "子智能体", "职责边界", "主要输出"],
        [
            ["1", "task-query-display-agent", "查询并展示任务，保留原始事实，不做设计决策。", "任务候选与完整任务快照"],
            ["2", "nx-template-agent", "选择/打开模板并确认 WorkPart，不做参数计算。", "模板来源与WorkPart确认"],
            ["3", "parameter-intelligence-agent", "抽证据、算参数、做写入计划，只读NX不写入。", "参数证据、公式结果、写入计划"],
            ["4", "nx-modeling-tc-agent", "执行已确认参数计划，回读、保存、TC回传。", "写入对比、保存结果、TC对象"],
            ["5", "drawing-update-agent", "选择图纸、打开视图、生成截图证据。", "图纸对象、视图清单、截图"],
            ["6", "dfmea-author-agent", "按受控模板生成并校验DFMEA。", "DFMEA工作簿与校验结果"],
            ["7", "design-report-author-agent", "按模板汇总证据生成设计说明书。", "设计说明书DOCX与证据清单"],
        ],
        widths=[1.2, 4.2, 7.0, 5.0],
    )

    doc.add_heading("4. 现有评测体系的使用方式", level=1)
    add_para(
        doc,
        r"评测体系以 C:\Users\ASUS\Documents\Codex\parts-agent-evaluation-workbench 为准。本报告不另造指标，只说明如何把它用于稳定性改进。",
    )
    add_table(
        doc,
        ["对象", "现有规模", "使用要求"],
        [
            ["任务", "7个任务", "每次问题必须先定位到对应任务，不允许只写“智能体效果差”。"],
            ["场景", "35个场景", "缺陷复现时优先使用已有场景；没有场景时补充新场景而不是改指标口径。"],
            ["检查点", "134个检查点", "失败必须落到检查点；硬性检查点失败时版本不得放行。"],
            ["记录", "58条记录", "同一修订前后必须比较记录，保留版本差异。"],
            ["权重", "功能正确性4、交互友好性2、异常处理2、数据一致性1、响应时效1", "优化优先级按权重和阻断等级排序。"],
        ],
        widths=[2.5, 3.8, 10.8],
    )
    add_para(
        doc,
        "建议每天只看三张清单：硬性检查点失败清单、分数下降最大的版本差异清单、同一场景重复失败清单。维护动作围绕这三张清单展开。",
    )

    doc.add_heading("5. 每个子智能体的架构与工作过程", level=1)
    for i, agent in enumerate(AGENTS, start=1):
        doc.add_heading(f"5.{i} {agent.title}", level=2)
        add_table(
            doc,
            ["项目", "内容"],
            [
                ["Agent文件", agent.agent_file],
                ["职责", agent.responsibility],
                ["Skill", agent.skills],
                ["工具", agent.tools],
                ["输出", agent.outputs],
                ["常见失败信号", agent.failure_signals],
                ["无代码改进抓手", agent.no_code_handles],
            ],
            widths=[3.2, 14.0],
        )
        add_diagram(
            doc,
            f"{agent.title}架构图",
            rendered[f"{agent.key}_arch"],
            f"图{2 + (i - 1) * 2 + 1}：{agent.title}的输入、Skill、工具、输出与改进反馈。",
        )
        add_diagram(
            doc,
            f"{agent.title}运行时序图",
            rendered[f"{agent.key}_workflow"],
            f"图{2 + (i - 1) * 2 + 2}：{agent.title}从接收任务到日志/评测回流的运行过程。",
        )

    doc.add_heading("6. 可执行的错误处理闭环", level=1)
    add_para(
        doc,
        "这一章是工程师处理单个失败的标准作业。每个失败都先形成一个失败包，再交给辅助智能体处理。辅助智能体只基于日志和评测数据给出结论，不允许凭经验猜测，也不允许直接要求改代码。",
    )
    doc.add_heading("6.1 第一步：把失败定位到评测工作台对象", level=2)
    add_para(
        doc,
        r"打开 C:\Users\ASUS\Documents\Codex\parts-agent-evaluation-workbench，或直接查看 app/src/data 下的 tasks.json、records.json、versions.json、rule-templates.json、point-categories.json。不要新建指标口径，先用已有 taskId、sceneId、pointId、ruleTemplateId、hardRequired、categoryId 和 weight 定位失败。",
    )
    add_table(
        doc,
        ["要填的字段", "从哪里取", "填写规则", "不合格写法"],
        [
            ["versionId", "records.json 或工作台版本选择", "记录失败版本和对比版本，例如 agent-0624-3 vs agent-0624-2。", "只写“最新版”。"],
            ["taskId / sceneId", "tasks.json / records.json", "必须定位到 7 个任务之一和具体场景，例如 task-dfmea / t6-001。", "只写“DFMEA不好”。"],
            ["pointId", "records.values 的键或场景检查点", "记录具体检查点，例如 t6-001-function-1。硬性检查点优先处理。", "只写“功能失败”。"],
            ["actual / expected", "records.values.actual、tasks.points.description", "actual 写实际记录值，expected 写检查点描述或规则阈值。", "用主观评价代替规则。"],
            ["sourceInput / note / response-time", "records.json", "保留原输入、备注、耗时和附件来源；耗时类问题必须记录秒数。", "只写“有点慢”。"],
        ],
        widths=[3.5, 5.0, 10.5, 5.5],
    )

    doc.add_heading("6.2 第二步：提取运行日志证据", level=2)
    add_para(
        doc,
        "运行日志由 ConversationTurnLogger 写入 JSONL 和 summary。已知 conversation_id 时，先用本地接口取日志状态；没有 conversation_id 时，从前端会话、测试记录附件或人工记录中补齐后再取证。",
    )
    add_code_block(
        doc,
        r"""
curl "http://127.0.0.1:8765/api/runtime/conversation-log?conversation_id=<conversation_id>"
""",
    )
    add_table(
        doc,
        ["日志字段", "必须提取的内容", "用途"],
        [
            ["log_path / summary_path", "JSONL和summary文件路径、exists、summary_exists。", "证明本轮有可追溯日志；没有日志不能关闭问题。"],
            ["turn", "turn_id、conversation_id、agent_id、mode、query、files、metadata。", "确认用户输入和入口智能体是否正确。"],
            ["mcp_tool / source", "本轮桥接来源和MCP工具名。", "确认是否进入了正确工具箱。"],
            ["server_message_type_counts", "session.tool.started/completed/result、session.failed等计数。", "判断是否中途失败、是否有工具调用缺失。"],
            ["tool_events", "工具开始、完成、结果摘要、artifact线索。", "重建时序，找出第一个错误工具或缺失工具。"],
            ["mcp_batches", "每次请求/响应摘要、terminal、status、returned_fields。", "定位工具返回是否失败、字段是否缺失。"],
            ["last_status / last_error", "失败码、失败消息、终态。", "区分业务失败、外部系统失败、协议/代码失败。"],
            ["connector_tools", "连接器工具清单和可用状态。", "排查工具不可用、工具定义未加载、权限缺失。"],
        ],
        widths=[4.5, 12.0, 8.0],
    )

    doc.add_heading("6.3 第三步：创建失败包", level=2)
    add_para(doc, "每个失败都先整理成下面这个 YAML。后续交给任何辅助智能体时，都只传这个失败包加原始日志，不传口头猜测。")
    add_code_block(
        doc,
        r"""
failure_packet:
  id: "FP-20260626-001"
  version_id: "agent-0624-3"
  task_id: "task-dfmea"
  scene_id: "t6-001"
  point_id: "t6-001-function-1"
  point_title: "存在疲劳断裂、孔磨损、孔变形等失效模式"
  category_id: "cat-function"
  hard_required: true
  rule_template_id: "rule-min-100"
  expected: "覆盖主要失效模式，且证据、风险分值、建议措施可读"
  actual: "actual=2；生成了DFMEA，但风险项覆盖和评分证据不足"
  source_input: "生成DFMEA表"
  response_time_sec: 111.677
  workbench_record_id: "rec-agent-0624-3-t6-001-r1"
  conversation_id: "<从测试记录或会话中补齐>"
  log_path: "<conversation jsonl path>"
  summary_path: "<conversation summary path>"
  artifacts:
    - "<DFMEA xlsx path or screenshot path>"
  business_impact: "任务6硬性功能点失败，影响DFMEA交付可靠性"
  current_owner: "维护工程师姓名"
""",
    )

    doc.add_heading("6.4 第四步：用辅助智能体逐个处理，不跳步", level=2)
    add_table(
        doc,
        ["顺序", "辅助智能体", "输入", "必须输出", "通过条件"],
        [
            ["1", "日志证据智能体", "failure_packet、JSONL全文、summary全文", "按时间排序的工具链；第一个异常点；缺失证据；artifact路径；不能判断的字段。", "所有结论都能指向日志字段或原始记录。"],
            ["2", "指标分析智能体", "failure_packet、tasks.json、records.json、versions.json、point-categories.json、rule-templates.json", "任务/场景/检查点定位；权重和硬性标记；当前版本与上个版本差异；受影响检查点清单。", "未新增指标；每个判断都有已有检查点或规则模板。"],
            ["3", "失败归因智能体", "日志证据输出、指标分析输出、7个Agent职责边界", "根因类别；归属子智能体；最小无代码修订对象；复测范围；是否需要升级代码缺陷。", "根因只能落到一个主因和最多两个次因，不能泛泛写多个可能。"],
            ["4", "指令/模板修订智能体", "归因结论、目标Agent/Skill文本、失败样本、通过样本", "可审查的修订文本；新增反例；输出格式要求；回滚点；预计影响场景。", "只修改Agent说明、Skill、模板、知识表、评测样本或SOP，不改.py/.cs/工具协议。"],
            ["5", "复测智能体", "修订说明、原失败场景、相邻场景、历史通过场景", "复测矩阵；每个检查点actual和note；版本对比；是否关闭建议。", "原失败场景连续两轮通过，且历史通过场景无硬性回退。"],
        ],
        widths=[1.6, 4.0, 7.2, 8.2, 4.5],
    )

    doc.add_heading("6.5 第五步：关闭或升级", level=2)
    add_table(
        doc,
        ["判定", "满足条件", "动作"],
        [
            ["关闭", "失败包完整；日志能解释问题；无代码修订已评审；原失败场景连续两轮通过；相邻场景和历史通过场景没有硬性回退。", "在维护记录中写入修订对象、版本、复测结果和关闭人。"],
            ["继续修订", "复测仍失败，但日志显示Agent指令、Skill流程、证据模板或输出格式仍有可修订空间。", "回到日志证据智能体，保留上一轮失败包和修订差异。"],
            ["升级代码缺陷", "工具入参正确但MCP/NX/TC/桥接持续返回失败；日志出现协议异常、工具不可用、返回字段缺失且Agent无法规避；同一工具在多个Agent复现。", "创建代码缺陷单，附JSONL、summary、最小复现输入、工具入参/出参、影响任务和检查点。"],
            ["暂缓", "外部系统不可用、业务资料缺失、用户确认信息缺失，导致无法判断。", "记录阻塞项和责任人；不得把暂缓问题计为修复完成。"],
        ],
        widths=[3.0, 15.0, 6.5],
    )

    doc.add_heading("7. 辅助智能体提示词模板", level=1)
    add_para(doc, "下面的提示词可以直接复制给相应辅助智能体。把尖括号内容替换为失败包、日志或文件内容。")
    add_table(
        doc,
        ["辅助智能体", "可直接使用的提示词", "输出格式"],
        [
            [
                "日志证据智能体",
                "请只依据以下 failure_packet、conversation JSONL 和 summary 分析，不要猜测。输出：1) 工具调用时间线；2) 第一个异常点；3) 缺失证据；4) artifact路径；5) 不能从日志判断的事项。禁止写业务建议。",
                "表格：序号、ts、agent_id、mcp_tool、工具事件、输入摘要、输出摘要、错误、证据字段。",
            ],
            [
                "指标分析智能体",
                "请读取 failure_packet 和评测工作台数据，只使用 tasks.json、records.json、versions.json、point-categories.json、rule-templates.json 中已有口径。输出失败检查点、权重、硬性标记、历史版本表现和优先级。",
                "表格：taskId、sceneId、pointId、category、weight、hardRequired、rule、current、baseline、priority。",
            ],
            [
                "失败归因智能体",
                "请根据日志证据、指标定位和7个子智能体职责边界做单主因归因。根因类别只能选择：路由错误、职责越界、证据缺失、公式/单位错误、模板/输出规则不足、外部系统异常、MCP/代码缺陷。",
                "结论卡：主因、次因、归属Agent、证据、为什么不是其他原因、最小修订对象、复测范围。",
            ],
            [
                "指令/模板修订智能体",
                "请只提出无代码修订。根据归因结论，给出需要改的Agent说明、Skill步骤、输出模板、知识表或评测样本。输出原文位置、建议文本、反例、影响范围、回滚方法。不得要求改.py/.cs/接口协议。",
                "修订清单：目标文件/章节、插入/替换文本、触发条件、禁止事项、回滚点、预期改善检查点。",
            ],
            [
                "复测智能体",
                "请根据修订清单制定复测。必须包含原失败场景、同任务相邻场景、至少一个历史通过场景；逐检查点记录actual、note、耗时和artifact；给出是否关闭建议。",
                "复测矩阵：versionId、taskId、sceneId、pointId、expected、actual、pass/fail、evidence、note。",
            ],
        ],
        widths=[3.8, 14.2, 6.5],
    )

    doc.add_heading("8. 无代码修订对象清单", level=1)
    add_para(
        doc,
        "这里的“无代码”不是不做改进，而是不改运行程序和工具协议。优先改可由维护工程师直接审查的文字、模板、知识和评测样本。所有修订都要能追溯到失败包。",
    )
    add_table(
        doc,
        ["修订对象", "适用问题", "具体改法", "验收方式"],
        [
            ["Agent说明", "路由错误、职责越界、工具误用、输出不完整", "增加触发条件、退出条件、禁止事项、必须回读项、交接输出格式；加入失败样本作为反例。", "原失败路由不再出现；Agent不调用越界工具；输出字段完整。"],
            ["Skill流程", "步骤缺失、顺序错误、缺少阻断门禁", "把隐含经验写成编号步骤，例如先查证据再计算、先回读再保存、校验失败不得交付。", "日志时间线符合Skill顺序；缺失证据时能阻断。"],
            ["输出模板", "字段缺失、格式不稳定、用户无法复核", "固定表头、字段顺序、单位、来源、置信度、待确认项；对错误和空结果给固定话术。", "同一场景多轮输出结构一致；检查点字段可直接打分。"],
            ["知识表/别名表", "模板选错、参数同义词误配、TC路径混乱", "增加模板优先级、路径别名、参数标准ID、中英文同义词；候选必须经工具回读确认。", "模板/参数匹配准确率提升；误配样本通过复测。"],
            ["公式和单位证据", "参数计算错误、单位换算错误", "公式来源必须记录；计算过程必须可读；单位换算形成显式检查项；边界值进入评测样本。", "任务3边界和缺证场景通过；任务4回读值与计划一致。"],
            ["交付模板", "DFMEA或报告质量低、章节/字段缺失", "固定必填字段、槽位、图片检查、validate门禁；缺失时输出待确认项而不是硬编。", "模板校验通过；交付物字段完整；缺证场景有清晰占位。"],
            ["评测样本", "问题已修但容易回归，或现有场景覆盖不足", "在现有任务下补充场景或记录，不改权重口径；新增样本需有sourceInput、expected、hardRequired说明。", "后续版本能复现旧失败，版本对比能看出改善。"],
            ["运行SOP", "工程师操作不一致、人工确认缺失", "写清谁确认、确认什么、何时停止、如何记录artifact；把口头经验写成检查清单。", "不同维护人员按同一SOP得到一致证据。"],
        ],
        widths=[3.2, 5.0, 10.5, 5.8],
    )

    doc.add_heading("9. 按子智能体落地的改进抓手", level=1)
    add_table(
        doc,
        ["子智能体", "优先看哪些日志/指标", "最常见可执行修订", "复测组合"],
        [
            ["task-query-display-agent", "records 的 sourceInput、响应时间、字段缺失note；日志中的 query_ipm_list/connect_qpp/query_ecr_list 输入输出和超时。", "固定任务列表表头：来源系统、任务编号、名称、状态、原始描述、附件；无结果必须给替代检索；多系统查询必须标注失败系统和已成功系统。", "t1-001关键词、t1-003多系统、t1-004无结果、一个历史通过场景。"],
            ["nx-template-agent", "Teamcenter候选路径、模板对象版本、nx_open_*结果、nx_get_work_part_info。", "模板选择规则增加优先级：任务部件族 > 指定模板编号 > TC分类 > 最近成功模板；必须回读WorkPart并输出部件号/名称/单位/状态。", "t2-001有效模板、t2-004高级模板、一个无效模板或路径别名场景。"],
            ["parameter-intelligence-agent", "公式来源、mysql_query结果、local_python_task校验、NX表达式resolve、单位字段。", "参数表固定列：参数ID、中文名、任务值、单位、公式、证据来源、NX表达式、模板当前值、是否需写入、待确认项；缺证不得计算。", "t3-001识别、t3-003边界值、缺证/单位换算场景、任务4回读联测。"],
            ["nx-modeling-tc-agent", "写入计划、NX update/batch/save、写后readback、Teamcenter upload/copy结果。", "执行前输出计划签名；执行后输出写前/写后对比；保存失败、回读不一致、TC上传失败均阻断，不允许交付成功结论。", "t4-001建模、参数变更场景、TC回传场景、一个无需写入场景。"],
            ["drawing-update-agent", "TC图纸候选、NX view list、switch/fit/create image、截图路径和文件存在性。", "图纸输出固定包含对象、版本、选择原因、视图清单、截图路径、样式是否恢复；空视图或截图缺失必须请求用户确认。", "t5所有图纸场景，至少覆盖图纸缺失、空视图、截图存在性。"],
            ["dfmea-author-agent", "template_list/inspect/select、calculate_risk、fill_template、validate_workbook、XLSX artifact。", "DFMEA风险项固定列：对象、功能、失效模式、影响、原因、现行控制、S/O/D、RPN、AP、建议措施、证据来源；每行必须调用或引用calculate_risk；validate失败不得交付。", "t6-001核心失效、t6-004字段完整、t6-005特殊材料、任务7报告一致性。"],
            ["design-report-author-agent", "template_inspect/start/collect/decrypt_image/answer/status/generate/artifact、图片和DOCX路径。", "章节证据矩阵固定化；图片必须decrypt/check后引用；缺失DFMEA或TC人员信息时写待确认项；generate后必须轮询status完成再返回artifact。", "t7-001完整数据、t7-002缺失数据、t7-003章节结构、t7-004一致性。"],
        ],
        widths=[4.0, 7.0, 9.5, 4.0],
    )

    doc.add_heading("10. 日常、每周、每版本的改进节奏", level=1)
    add_table(
        doc,
        ["频率", "工程师实际动作", "交给哪个辅助智能体", "必须产物", "放行门槛"],
        [
            ["每天", "在工作台看硬性检查点失败、分数下降、重复失败；每个失败建立failure_packet。", "指标分析智能体", "当日失败队列，按任务、严重度、版本排序。", "所有硬性失败都有owner和下一步。"],
            ["每个失败", "取summary和JSONL，补齐日志路径、artifact、sourceInput、actual/expected。", "日志证据智能体", "证据时间线和第一个异常点。", "没有日志不允许写修复结论。"],
            ["每个修订", "只改Agent/Skill/模板/知识/样本/SOP；记录原文和新文本。", "指令/模板修订智能体", "无代码修订清单和回滚点。", "人工评审确认不越界、不编造事实。"],
            ["每周", "汇总同类根因，合并重复规则，删除互相冲突的说明。", "失败归因智能体", "高频根因榜和下一周修订优先级。", "优先处理高权重、硬性、重复出现问题。"],
            ["每版本", "跑原失败场景、相邻场景、历史通过场景，更新records和versions。", "复测智能体", "版本对比报告和关闭建议。", "无新增硬性失败；原失败连续两轮通过。"],
        ],
        widths=[2.4, 8.0, 4.6, 6.0, 3.5],
    )

    doc.add_heading("11. 典型问题的处理样例", level=1)
    doc.add_heading("11.1 DFMEA任务低分的处理样例", level=2)
    add_numbered(
        doc,
        [
            "定位：在工作台找到 task-dfmea 下 t6-001、t6-004、t6-005 的失败或低分记录，优先处理 hardRequired=true 且 cat-function/cat-consistency 权重高的检查点。",
            "取证：从日志中提取 dfmea_template_list、inspect、select、calculate_risk、fill_template、validate_workbook 的调用顺序，确认是否每条风险都经过风险计算，是否校验通过后才返回XLSX。",
            "归因：如果模板和XLSX存在但分值低，通常不是代码问题，而是风险项字段、证据来源、S/O/D依据、特殊材料失效模式不足。",
            "无代码修订：在 dfmea-risk-authoring 中增加风险项必填列和证据来源要求；在 dfmea-workbook-generation 中增加 validate_workbook失败阻断；在输出模板中要求列出每行S/O/D/RPN/AP依据。",
            "补样本：在任务6下补充特殊材料、缺少边界图、模板字段缺失、风险计算偏差四类样本；不改类别权重。",
            "复测：先跑 t6-001 和 t6-004，再跑 t6-005，最后跑任务7中依赖DFMEA的报告一致性场景。全部通过后关闭。",
        ],
    )
    doc.add_heading("11.2 参数计算错误的处理样例", level=2)
    add_numbered(
        doc,
        [
            "定位：找 task-parameters 的失败检查点，尤其是边界值、单位换算、参数映射和缺证场景。",
            "取证：查看 mysql_query 公式来源、local_python_task脚本校验结果、nx_find_params/nx_resolve_parameter 对应的NX表达式和单位。",
            "归因：如果公式来源缺失或单位未显式换算，归因到Skill/输出模板；如果工具返回表达式错误且输入正确，才升级工具缺陷。",
            "无代码修订：参数表必须输出公式来源、计算过程、单位、NX表达式、模板当前值、任务计算值、是否写入；缺证时输出待确认项，不得继续计算。",
            "复测：任务3原失败场景、边界场景、缺证场景，以及任务4写后回读场景。",
        ],
    )
    doc.add_heading("11.3 任务查询不稳定的处理样例", level=2)
    add_numbered(
        doc,
        [
            "定位：看 task-retrieval 的 t1-001、t1-003、t1-004、t1-005；多系统查询和无结果查询优先。",
            "取证：从日志中区分 IPM、QPP、ECR 哪个系统超时，是否已有部分结果，输出是否保留原始任务字段。",
            "归因：单个外部系统超时不算Agent错误；但未标注已成功系统、未保留原文、未给替代建议属于Agent输出规则问题。",
            "无代码修订：输出固定为来源系统、任务编号、任务名称、状态、原始描述、附件线索、下一步选择；超时系统单独列出，不影响已返回系统展示。",
            "复测：关键词、多系统联合、无结果、分页四类场景各跑一轮；同一输入至少两轮检查字段稳定性。",
        ],
    )

    doc.add_heading("12. 不允许关闭的问题", level=1)
    add_table(
        doc,
        ["情况", "为什么不能关闭", "正确动作"],
        [
            ["只有主观描述，没有失败包", "无法定位到任务、场景和检查点，不能评估修复效果。", "补齐failure_packet。"],
            ["只有评测记录，没有运行日志", "不知道是Agent、工具、外部系统还是人工操作导致。", "补取JSONL和summary；取不到则标为证据不足。"],
            ["只改了说明，没有复测", "无法证明稳定性和正确率提升。", "交给复测智能体跑原失败和相邻场景。"],
            ["只跑成功样例", "不能防止历史通过场景回退。", "至少加一个历史通过场景和一个边界场景。"],
            ["代码异常被包装成提示词问题", "无代码修订无法解决真实工具缺陷。", "升级代码缺陷单，附最小复现和工具入参出参。"],
            ["外部系统不可用时强行改Agent", "会把环境问题误修成错误规则。", "标记外部依赖不可用，复测时注明环境状态。"],
        ],
        widths=[5.0, 11.0, 8.5],
    )

    doc.add_heading("13. 交付与维护记录", level=1)
    add_para(doc, "本版文档中的所有架构图和时序图均由 Mermaid 源渲染为 PNG 后嵌入 Word，Word 中不放 Mermaid 源码块。")
    add_para(doc, "建议在每次修订后新增一条维护记录，至少包含：失败编号、相关日志、评测检查点、修订对象、修订文本、复测结果、是否关闭。")
    add_para(doc, "生成文件位置：")
    add_hyperlink_like_path(doc, DOCX_PATH)

    doc.save(DOCX_PATH)


def build_architecture_doc(rendered: dict[str, Path]) -> None:
    doc = setup_document()
    add_title_page(
        doc,
        title="零部件设计智能体软件模型架构说明书",
        subtitle="项目验收技术报告用 | 软件模型、运行链路、7个子智能体架构与过程说明",
        version="Mermaid截图版",
        audience="项目验收组 / 技术评审人员 / 系统架构人员",
        basis="项目代码、Agent定义、MCP工具边界、运行日志实现",
        doc_type="技术架构说明",
    )

    doc.add_heading("1. 文档定位", level=1)
    add_para(
        doc,
        "本文用于项目验收技术报告，重点说明 mc-design 项目中设计智能体的软件模型架构、运行链路、7个任务子智能体的职责边界、工具依赖、输入输出和运行过程。本文不展开持续改进SOP；持续改进方法见配套文档《零部件设计智能体持续改进操作说明书》。",
    )
    add_bullets(
        doc,
        [
            "验收关注点一：主智能体是否只做任务编排和人机确认，不越权调用底层设计工具。",
            "验收关注点二：7个子智能体是否各自承担独立任务边界，工具权限与职责一致。",
            "验收关注点三：每个关键步骤是否有可回溯的工具调用、输出物和日志证据。",
            "验收关注点四：NX、Teamcenter、MySQL、DFMEA、设计报告等外部能力是否通过本地 MCP 工具箱受控接入。",
        ],
    )

    doc.add_heading("2. 软件模型总体架构", level=1)
    add_para(
        doc,
        "系统采用“云端/服务侧调度 + 本地客户端桥接 + MCP工具箱 + 业务系统”的软件模型。AI Service 通过 RuntimeBridgeManager 将 turn-stream-v1 调用桥接到本地 ClientApp；ClientApp 负责暴露 /mcp/<toolbox> 工具入口，并通过 TJUAE LiveCoordinator 同步会话、插件和工具定义。design-agent 是默认入口，只负责路由和必要确认，具体任务由7个子智能体完成。",
    )
    add_diagram(doc, "总体软件模型架构图", rendered["overall_arch"], "图1：设计智能体总体软件模型、7个子智能体、MCP工具箱和企业系统关系。")

    doc.add_heading("3. 核心组件说明", level=1)
    add_table(
        doc,
        ["组件", "项目位置/来源", "验收时应确认的内容"],
        [
            ["插件默认入口", r"mc-design-nx/client/plugins/mc-design/.tjuae-plugin/plugin.json", "默认 agent 为 mc-design:design-agent；MCP servers 指向 connector、mysql、teamcenter、design_report、dfmea、local_file、nx 等工具箱。"],
            ["主智能体", "design-agent.md", "只调用7个子Agent和AskUserQuestion；不直接操作NX、TC、MySQL、DFMEA或报告工具。"],
            ["子智能体定义", "*.agent.md", "每个子智能体有明确Skill、工具白名单和禁止事项。"],
            ["本地客户端", "mc_design_client/app.py", "负责ClientApp、工具定义加载、桥接请求和call_agent入口。"],
            ["运行桥接", "mc-design-ai-service/beya_mcp/runtime_bridge.py", "RuntimeBridgeManager 通过 turn-stream-v1 协议把服务侧请求路由到本地运行时。"],
            ["运行日志", "runtime/conversation_log.py", "ConversationTurnLogger 记录 JSONL 和 summary，包括tool_events、mcp_batches、last_error等验收证据。"],
        ],
        widths=[4.0, 8.0, 12.5],
    )

    doc.add_heading("4. 7个子智能体边界总览", level=1)
    add_table(
        doc,
        ["任务", "子智能体", "职责边界", "主要工具/系统", "主要交付"],
        [
            ["1", "task-query-display-agent", "任务抓取与展示，保留原始任务事实。", "IPM/QPP/ECR connector、local_python_task", "任务候选、任务快照"],
            ["2", "nx-template-agent", "模板选择、打开和WorkPart确认。", "Teamcenter folder/classification、NX open/create/read", "模板来源、WorkPart确认"],
            ["3", "parameter-intelligence-agent", "证据抽取、公式计算、参数映射、写入计划；只读NX。", "MySQL、NX read/find/resolve、local_python_task", "参数证据、计算结果、写入计划"],
            ["4", "nx-modeling-tc-agent", "执行已确认参数计划、回读、保存、TC回传。", "NX update/save、Teamcenter upload/copy/export", "写入前后对比、TC对象"],
            ["5", "drawing-update-agent", "图纸选择、视图切换、截图证据。", "Teamcenter children/export、NX drawing/view/image", "图纸对象、视图清单、截图"],
            ["6", "dfmea-author-agent", "受控模板DFMEA生成和校验。", "DFMEA template/risk/fill/validate、TC export", "DFMEA XLSX、校验结果"],
            ["7", "design-report-author-agent", "证据收集和设计说明书DOCX生成。", "design_report工具、local_file", "设计说明书、证据清单"],
        ],
        widths=[1.5, 4.2, 7.5, 6.0, 5.3],
    )

    doc.add_heading("5. 子智能体架构与运行过程", level=1)
    for i, agent in enumerate(AGENTS, start=1):
        doc.add_heading(f"5.{i} {agent.title}", level=2)
        add_para(doc, agent.responsibility)
        add_table(
            doc,
            ["架构项", "说明"],
            [
                ["Agent文件", agent.agent_file],
                ["Skill", agent.skills],
                ["工具白名单", agent.tools],
                ["输出物", agent.outputs],
                ["验收关注点", agent.failure_signals],
            ],
            widths=[3.4, 21.0],
        )
        add_diagram(
            doc,
            f"{agent.title}架构图",
            rendered[f"{agent.key}_arch"],
            f"图{1 + (i - 1) * 2 + 1}：{agent.title}的软件结构、输入输出、工具依赖和日志回流。",
        )
        add_diagram(
            doc,
            f"{agent.title}运行时序图",
            rendered[f"{agent.key}_workflow"],
            f"图{1 + (i - 1) * 2 + 2}：{agent.title}的运行时序、工具调用和证据记录过程。",
        )

    doc.add_heading("6. 项目验收建议检查项", level=1)
    add_table(
        doc,
        ["检查项", "验收方法", "通过标准"],
        [
            ["入口和路由", "检查plugin.json默认agent和design-agent工具列表。", "默认入口为design-agent，且只能分派7个子智能体和AskUserQuestion。"],
            ["职责隔离", "逐个检查Agent文件中的工具白名单和禁止事项。", "子智能体工具权限与职责一致；参数计算Agent不写NX，建模Agent不重新计算参数。"],
            ["运行证据", "抽查conversation JSONL和summary。", "能看到turn、mcp_tool、tool_events、mcp_batches、last_error和artifact路径。"],
            ["交付物证据", "抽查模型、截图、DFMEA、设计说明书等输出。", "交付物能回溯到对应工具调用、用户确认和任务场景。"],
            ["异常处理", "使用评测工作台中的异常场景。", "失败时不空白、不伪造结果，能给出阻断原因和下一步确认项。"],
        ],
        widths=[4.5, 10.0, 10.0],
    )

    doc.add_heading("7. 文件位置", level=1)
    add_para(doc, "本文输出文件：")
    add_hyperlink_like_path(doc, ARCH_DOCX_PATH)
    doc.save(ARCH_DOCX_PATH)


def build_operations_doc(rendered: dict[str, Path]) -> None:
    doc = setup_document()
    add_title_page(
        doc,
        title="零部件设计智能体持续改进操作说明书",
        subtitle="技术维护与迭代升级用 | 日志取证、评测定位、辅助智能体归因、无代码修订与复测门禁",
        version="Mermaid截图版",
        audience="技术维护人员 / 智能体迭代升级人员 / 评测负责人",
        basis="现有评测工作台、运行日志、Agent/Skill/模板维护流程",
        doc_type="维护操作手册",
    )

    doc.add_heading("1. 使用边界", level=1)
    add_para(
        doc,
        "本文用于后续维护和迭代升级。常规改进不是直接改代码，而是由辅助智能体基于运行日志和评测指标进行证据提取、指标定位、失败归因、无代码修订建议和复测。只有日志明确证明 MCP、NX、Teamcenter、桥接协议或程序实现错误时，才升级为代码缺陷。",
    )
    add_diagram(doc, "持续改进闭环图", rendered["improvement_loop"], "图1：日志证据、指标分析、失败归因、指令修订和复测构成的闭环。")

    doc.add_heading("2. 维护数据源", level=1)
    add_table(
        doc,
        ["数据源", "路径/入口", "维护时使用方式"],
        [
            ["评测工作台", r"C:\Users\ASUS\Documents\Codex\parts-agent-evaluation-workbench", "管理任务、场景、检查点、版本和测试记录；不另造指标。"],
            ["任务/检查点", r"app/src/data/tasks.json", "读取 taskId、sceneId、pointId、description、hardRequired、ruleTemplateId。"],
            ["测试记录", r"app/src/data/records.json", "读取 versionId、sourceInput、actual、note、response time、attachments。"],
            ["版本信息", r"app/src/data/versions.json", "做当前版本和基线版本对比。"],
            ["规则和分类", r"rule-templates.json / point-categories.json", "引用已有规则、权重和分类，不新增口径。"],
            ["运行日志", "conversation JSONL + .summary.json", "读取turn、tool_events、mcp_batches、last_error、artifact路径。"],
            ["日志接口", r"http://127.0.0.1:8765/api/runtime/conversation-log?conversation_id=<id>", "已知conversation_id时先取日志状态和summary。"],
        ],
        widths=[4.0, 8.5, 12.0],
    )

    doc.add_heading("3. 单个失败的标准处理流程", level=1)
    add_numbered(
        doc,
        [
            "在评测工作台定位失败：写清 versionId、taskId、sceneId、pointId、hardRequired、ruleTemplateId、actual、expected。",
            "取运行日志：获取 conversation JSONL 和 summary，保留 tool_events、mcp_batches、last_error、artifact 路径。",
            "准备失败处理材料：把评测定位、日志证据、用户输入、输出物路径、环境状态和业务影响收齐，格式由团队统一工具或维护流程确定。",
            "交给日志证据智能体：只让它还原工具调用链和第一个异常点，不让它给业务建议。",
            "交给指标分析智能体：只使用工作台已有任务、场景、检查点、规则、权重和版本记录。",
            "交给失败归因智能体：输出单主因、归属子智能体、最小修订对象和复测范围。",
            "交给指令/模板修订智能体：只提出无代码修订，包括Agent说明、Skill、输出模板、知识表、评测样本或SOP。",
            "人工评审修订：确认不越权、不编造事实、不绕过工具校验、不破坏历史通过场景。",
            "交给复测智能体：复测原失败场景、相邻场景和至少一个历史通过场景。",
            "关闭或升级：连续两轮通过且无新增硬性失败才关闭；证据指向工具/代码时升级缺陷单。",
        ],
    )

    doc.add_heading("4. 失败处理前准备材料清单", level=1)
    add_para(doc, "每个失败交给辅助智能体前，只要求准备完整材料，不强制工程师手工整理成某一种文本格式。团队后续可以在工作台或表单中统一承载这些字段。")
    add_table(
        doc,
        ["材料类型", "必须准备的内容", "来源", "用途"],
        [
            ["评测定位材料", "失败版本、任务、场景、检查点、检查点标题、分类、权重、是否硬性、规则模板、期望结果和实际结果。", "工作台；tasks.json；records.json；rule-templates.json；point-categories.json", "让指标分析智能体准确定位失败，不重新发明指标。"],
            ["原始输入材料", "用户原始输入、测试轮次、测试日期、测试人员、来源文件、相关附件说明。", "records.json；测试记录附件；人工测试记录", "保证后续复测能复现同一输入。"],
            ["运行日志材料", "conversation_id、JSONL日志、summary日志、tool_events、mcp_batches、last_error、connector_tools、artifact路径。", "conversation-log接口；客户端logs/conversations目录", "让日志证据智能体还原工具调用链和第一个异常点。"],
            ["输出物材料", "模型、截图、DFMEA工作簿、设计说明书、TC对象、报告artifact、文件存在性检查结果。", "工具返回结果；local_file；交付目录；测试附件", "证明失败发生在哪个交付环节，避免只凭回复文字判断。"],
            ["环境状态材料", "NX、Teamcenter、IPM/QPP/ECR、MySQL、DFMEA服务、设计报告服务是否可用；是否存在超时或权限异常。", "工具日志；服务状态；人工确认", "区分智能体问题、外部系统问题和代码/工具问题。"],
            ["影响和责任材料", "业务影响、阻断程度、责任维护人、期望修复版本、是否需要升级代码缺陷。", "维护记录；项目计划；人工确认", "决定优先级、关闭条件和升级路径。"],
            ["复测基线材料", "原失败场景、相邻场景、至少一个历史通过场景、基线版本记录。", "工作台records/versions；历史测试报告", "防止只修单点问题导致历史场景回退。"],
        ],
        widths=[3.8, 7.8, 6.2, 6.7],
    )

    doc.add_heading("5. 辅助智能体使用模板", level=1)
    add_table(
        doc,
        ["辅助智能体", "输入", "可直接复制的要求", "必须输出"],
        [
            ["日志证据智能体", "运行日志材料、原始输入材料、输出物材料", "只依据日志分析，不猜测。输出工具调用时间线、第一个异常点、缺失证据、artifact路径和不能判断事项。", "时间线表、异常点、证据字段、不可判断清单。"],
            ["指标分析智能体", "评测定位材料、版本记录、规则和分类数据", "只使用评测工作台已有口径，不新建指标。输出检查点、权重、硬性标记、版本差异和优先级。", "任务-场景-检查点矩阵、优先级。"],
            ["失败归因智能体", "日志证据、指标定位、Agent边界", "根因类别只能选路由错误、职责越界、证据缺失、公式/单位错误、模板/输出规则不足、外部系统异常、MCP/代码缺陷。", "单主因、次因、归属Agent、最小修订对象、复测范围。"],
            ["指令/模板修订智能体", "归因结论、目标Agent/Skill、失败样本、通过样本", "只提出无代码修订。给出原文位置、建议文本、反例、影响范围和回滚方法。", "修订清单、插入/替换文本、回滚点。"],
            ["复测智能体", "修订清单、原失败场景、相邻场景、历史通过场景", "逐检查点记录 actual、note、耗时和artifact；给出是否关闭建议。", "复测矩阵、版本对比、关闭建议。"],
        ],
        widths=[4.0, 5.5, 10.0, 5.0],
    )

    doc.add_heading("6. 无代码修订对象", level=1)
    add_table(
        doc,
        ["对象", "适用失败", "具体动作", "验收方式"],
        [
            ["Agent说明", "路由错误、职责越界、工具误用", "增加触发条件、退出条件、禁止事项、必须回读项、交接输出格式。", "日志显示调用正确Agent和工具，输出字段完整。"],
            ["Skill流程", "步骤缺失、顺序错误、门禁不足", "把隐含经验写成编号步骤，例如先证据后计算、先回读后保存、校验失败不得交付。", "工具调用顺序符合Skill，缺证时能阻断。"],
            ["输出模板", "字段缺失、格式不稳定", "固定表头、单位、来源、置信度、待确认项、错误话术。", "同一场景多轮输出结构一致。"],
            ["知识/别名表", "模板误选、参数误配、路径混乱", "增加模板优先级、TC路径别名、参数标准ID、中英文同义词；候选必须经工具回读确认。", "误配样本通过，历史通过样本不回退。"],
            ["公式和单位证据", "参数计算错误", "公式来源必须记录，单位换算显式输出，边界值进入评测样本。", "任务3边界/缺证场景和任务4回读场景通过。"],
            ["交付模板", "DFMEA或报告质量低", "固定必填字段、槽位、图片检查、validate门禁；缺失时写待确认项。", "交付物校验通过，缺证场景不伪造。"],
            ["评测样本", "旧问题容易回归", "在现有任务下补充场景或记录，不改权重口径。", "版本对比能复现旧失败并证明改善。"],
        ],
        widths=[3.5, 5.5, 10.5, 5.0],
    )

    doc.add_heading("7. 按子智能体的改进抓手", level=1)
    add_table(
        doc,
        ["子智能体", "先看日志/指标", "无代码改法", "复测组合"],
        [
            ["task-query-display-agent", "sourceInput、响应时间、query_ipm_list/connect_qpp/query_ecr_list输入输出和超时。", "固定任务列表字段；超时系统单列；无结果给替代检索；保留原始任务描述。", "t1-001、t1-003、t1-004、t1-005。"],
            ["nx-template-agent", "TC候选路径、模板版本、nx_open结果、WorkPart回读。", "模板选择优先级和反例；强制回读部件号/名称/单位/状态。", "t2-001、t2-004、无效模板场景。"],
            ["parameter-intelligence-agent", "公式来源、脚本校验、NX表达式resolve、单位字段。", "参数表固定列；缺证不得计算；单位换算显式输出。", "t3-001、t3-003、缺证场景、任务4联测。"],
            ["nx-modeling-tc-agent", "计划签名、NX写入/save/readback、TC upload结果。", "写前计划、写后对比、保存失败阻断、TC回传清单模板化。", "t4-001、参数变更、TC回传、无需写入。"],
            ["drawing-update-agent", "TC图纸候选、view list、create image、截图文件存在性。", "图纸选择原因、视图清单、截图路径、样式恢复固定输出；空视图阻断。", "任务5图纸缺失、空视图、截图场景。"],
            ["dfmea-author-agent", "template/inspect/select、calculate_risk、fill、validate、XLSX artifact。", "风险项必填列；每行S/O/D/RPN/AP依据；validate失败不得交付。", "t6-001、t6-004、t6-005、任务7一致性。"],
            ["design-report-author-agent", "template_inspect、collect、decrypt_image、status、artifact。", "章节证据矩阵；图片检查；缺失DFMEA/人员信息写待确认；status完成后返回。", "t7-001、t7-002、t7-003、t7-004。"],
        ],
        widths=[4.0, 7.0, 9.0, 4.5],
    )

    doc.add_heading("8. 关闭门禁", level=1)
    add_table(
        doc,
        ["情况", "是否允许关闭", "处理动作"],
        [
            ["只有主观描述，没有失败处理材料", "不允许", "补齐评测定位材料、运行日志材料、输出物材料和业务影响。"],
            ["只有评测记录，没有运行日志", "不允许", "补取JSONL和summary；取不到则标为证据不足。"],
            ["只改了提示词/模板，没有复测", "不允许", "复测原失败、相邻场景和历史通过场景。"],
            ["原失败通过但历史通过场景回退", "不允许", "回滚或继续修订，直到无硬性回退。"],
            ["外部系统不可用", "暂缓", "记录外部依赖状态，不把环境问题计为Agent修复。"],
            ["工具入参正确但MCP/NX/TC/桥接持续失败", "升级", "创建代码缺陷单，附日志、最小复现和工具入参出参。"],
            ["原失败连续两轮通过且无新增硬性失败", "允许", "归档失败处理材料、修订清单、复测矩阵和关闭人。"],
        ],
        widths=[7.0, 3.5, 14.0],
    )

    doc.add_heading("9. 文件位置", level=1)
    add_para(doc, "本文输出文件：")
    add_hyperlink_like_path(doc, OPS_DOCX_PATH)
    doc.save(OPS_DOCX_PATH)


def audit_docx(path: Path, *, min_images: int, required_terms: list[str], forbidden_terms: list[str] | None = None) -> None:
    import zipfile

    with zipfile.ZipFile(path) as zf:
        media = [name for name in zf.namelist() if name.startswith("word/media/")]
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    if len(media) < min_images:
        raise RuntimeError(f"{path.name}: expected at least {min_images} embedded images, got {len(media)}")
    if "```mermaid" in document_xml or "sequenceDiagram" in document_xml or "flowchart LR" in document_xml or "flowchart TB" in document_xml:
        raise RuntimeError(f"{path.name}: DOCX still contains Mermaid source text")
    text = "\n".join(par.text for par in Document(path).paragraphs)
    missing = [term for term in required_terms if term not in text]
    if missing:
        raise RuntimeError(f"{path.name}: missing required terms: {missing}")
    forbidden_found = [term for term in (forbidden_terms or []) if term in text]
    if forbidden_found:
        raise RuntimeError(f"{path.name}: forbidden terms found: {forbidden_found}")


def main() -> None:
    diagrams = {
        "overall_arch": OVERALL_ARCH,
        "improvement_loop": IMPROVEMENT_LOOP,
    }
    for agent in AGENTS:
        diagrams[f"{agent.key}_arch"] = agent.architecture
        diagrams[f"{agent.key}_workflow"] = agent.workflow

    write_mermaid_sources(diagrams)
    rendered = render_diagrams(diagrams)
    build_architecture_doc(rendered)
    build_operations_doc(rendered)
    audit_docx(
        ARCH_DOCX_PATH,
        min_images=15,
        required_terms=["软件模型总体架构", "7个子智能体边界总览", "项目验收建议检查项"],
    )
    audit_docx(
        OPS_DOCX_PATH,
        min_images=1,
        required_terms=["失败处理前准备材料清单", "辅助智能体使用模板", "无代码修订对象", "关闭门禁"],
        forbidden_terms=["failure_packet", "结构化 YAML"],
    )
    print(ARCH_DOCX_PATH)
    print(OPS_DOCX_PATH)


if __name__ == "__main__":
    main()
